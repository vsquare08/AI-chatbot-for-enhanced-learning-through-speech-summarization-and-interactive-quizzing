from flask import Flask,render_template,redirect,request,url_for, send_file, session
import mysql.connector
import pandas as pd
import numpy as np
from sklearn.metrics import mean_absolute_error
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler
from tensorflow.keras.models import Sequential
from tensorflow.keras.layers import Dense, LSTM, Bidirectional, SimpleRNN
from tensorflow.keras.optimizers import Adam
from sklearn.preprocessing import LabelEncoder
import joblib
from tensorflow.keras.models import load_model
from flask import Flask, render_template, request
from transformers import AutoTokenizer, AutoModelForSeq2SeqLM
from docx import Document
from PyPDF2 import PdfReader
import os
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from transformers import AutoTokenizer, BartForConditionalGeneration


app = Flask(__name__)
app.secret_key = 'your_secret_key'


mydb = mysql.connector.connect(
    host="localhost",
    user="root",
    password="",
    port="3306",
    database='AIteacher'
)

mycursor = mydb.cursor()

def executionquery(query,values):
    mycursor.execute(query,values)
    mydb.commit()
    return

def retrivequery1(query,values):
    mycursor.execute(query,values)
    data = mycursor.fetchall()
    return data

def retrivequery2(query):
    mycursor.execute(query)
    data = mycursor.fetchall()
    return data


@app.route('/')
def index():
    return render_template('index.html')

@app.route('/about')
def about():
    return render_template('about.html')


@app.route('/register', methods=["GET", "POST"])
def register():
    if request.method == "POST":
        name = request.form['name']
        email = request.form['email']
        password = request.form['password']
        c_password = request.form['c_password']
        if password == c_password:
            query = "SELECT UPPER(email) FROM users"
            email_data = retrivequery2(query)
            email_data_list = []
            for i in email_data:
                email_data_list.append(i[0])
            if email.upper() not in email_data_list:
                query = "INSERT INTO users (name, email, password) VALUES (%s, %s, %s)"
                values = (name, email, password)
                executionquery(query, values)
                return render_template('login.html', message="Successfully Registered!")
            return render_template('register.html', message="This email ID is already exists!")
        return render_template('register.html', message="Conform password is not match!")
    return render_template('register.html')


@app.route('/login', methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form['email']
        password = request.form['password']
        
        query = "SELECT UPPER(email) FROM users"
        email_data = retrivequery2(query)
        email_data_list = []
        for i in email_data:
            email_data_list.append(i[0])

        if email.upper() in email_data_list:
            query = "SELECT UPPER(password) FROM users WHERE email = %s"
            values = (email,)
            password__data = retrivequery1(query, values)
            if password.upper() == password__data[0][0]:
                global user_email
                user_email = email

                return redirect("/home")
            return render_template('login.html', message= "Invalid Password!!")
        return render_template('login.html', message= "This email ID does not exist!")
    return render_template('login.html')


@app.route('/home')
def home():
    return render_template('home.html')





df = pd.read_csv('self_harm_trends.csv')
df['timestamp'] = pd.to_datetime(df['timestamp'])
original_columns = df.select_dtypes(include='object').columns

# Initialize LabelEncoder
label_encoders = {}

# Apply LabelEncoder to each categorical variable
for col in original_columns:
    label_encoders[col] = LabelEncoder()
    df[col] = label_encoders[col].fit_transform(df[col])

# Drop timestamp for the prediction task
df.drop('timestamp', axis=1, inplace=True)

# Define features and target variable
X = df.drop('self_harm_severity', axis=1)
y = df['self_harm_severity']


# Split data
X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)

# Standardize the data
scaler = StandardScaler()
X_train = scaler.fit_transform(X_train)
X_test = scaler.transform(X_test)

# Reshape the input data to 3D (samples, timesteps, features)
X_train_reshaped = np.expand_dims(X_train, axis=1)
X_test_reshaped = np.expand_dims(X_test, axis=1)

def compile_and_train(model, X_train, y_train, X_test, y_test, epochs=10, batch_size=32):
    model.compile(optimizer=Adam(learning_rate=0.001), loss='mean_squared_error', metrics=['mae'])
    history = model.fit(X_train, y_train, validation_data=(X_test, y_test), epochs=epochs, batch_size=batch_size, verbose=1)
    return model, history


# Load the pretrained model and tokenizer
model_path = "summarization_model"
tokenizer = AutoTokenizer.from_pretrained(model_path)
model = AutoModelForSeq2SeqLM.from_pretrained(model_path)

@app.route('/algorithm_selection', methods=["GET", "POST"])
def algorithm_selection():
    summary = None  # Default value for the summary
    if request.method == "POST":
        uploaded_file = request.files.get("file")
        summary_ratio = request.form.get("summary_ratio", "30")  # Default to 30%
        text = ""

        # Extract text from uploaded file
        if uploaded_file:
            file_extension = os.path.splitext(uploaded_file.filename)[1].lower()
            if file_extension == ".docx":
                doc = Document(uploaded_file)
                text = "\n".join([para.text for para in doc.paragraphs])
            elif file_extension == ".pdf":
                reader = PdfReader(uploaded_file)
                text = "\n".join([page.extract_text() for page in reader.pages])
            else:
                return render_template(
                    'algorithm_selection.html',
                    error="Unsupported file format. Please upload a .docx or .pdf file.",
                    summary=None
                )

        if not text.strip():
            return render_template(
                'algorithm_selection.html',
                error="The uploaded file contains no text.",
                summary=None
            )

        # Process the summary ratio
        try:
            summary_ratio = float(summary_ratio) / 100  # Convert to a decimal
            if summary_ratio <= 0 or summary_ratio > 1:
                raise ValueError("Summary ratio must be between 0 and 100.")
        except ValueError as e:
            return render_template('algorithm_selection.html', error=str(e), summary=None)

        # Tokenize and summarize the text
        try:
            inputs = tokenizer(text, return_tensors="pt", max_length=1024, truncation=True)
            input_length = len(tokenizer.tokenize(text))
            max_length = int(input_length * summary_ratio)
            min_length = max(30, int(max_length * 0.5))

            summary_ids = model.generate(
                inputs["input_ids"],
                max_length=max_length,
                min_length=min_length,
                length_penalty=2.0,
                num_beams=4,
                early_stopping=True
            )

            summary = tokenizer.decode(summary_ids[0], skip_special_tokens=True)

        except Exception as e:
            return render_template('algorithm_selection.html', error=str(e), summary=None)

    return render_template('algorithm_selection.html', summary=summary, error=None)

Load the pretrained BART model and tokenizer


# Load the tokenizer and the BART model for summarization
# model_path = "BARTsummarization_model"
# tokenizer = AutoTokenizer.from_pretrained(model_path)
# model = BartForConditionalGeneration.from_pretrained(model_path)

# @app.route('/algorithm_selection', methods=["GET", "POST"])
# def algorithm_selection():
#     summary = None  # Default value for the summary
#     if request.method == "POST":
#         uploaded_file = request.files.get("file")
#         summary_ratio = request.form.get("summary_ratio", "30")  # Default to 30%
#         text = ""

#         # Extract text from uploaded file
#         if uploaded_file:
#             file_extension = os.path.splitext(uploaded_file.filename)[1].lower()
#             if file_extension == ".docx":
#                 doc = Document(uploaded_file)
#                 text = "\n".join([para.text for para in doc.paragraphs])
#             elif file_extension == ".pdf":
#                 reader = PdfReader(uploaded_file)
#                 text = "\n".join([page.extract_text() for page in reader.pages])
#             else:
#                 return render_template(
#                     'algorithm_selection.html',
#                     error="Unsupported file format. Please upload a .docx or .pdf file.",
#                     summary=None
#                 )

#         if not text.strip():
#             return render_template(
#                 'algorithm_selection.html',
#                 error="The uploaded file contains no text.",
#                 summary=None
#             )

#         # Tokenize the input text
#         inputs = tokenizer(text, return_tensors="pt", max_length=1024, truncation=True)

#         # Generate the summary
#         summary_ids = model.generate(
#             inputs["input_ids"],
#             max_length=150,  # Adjust max_length as needed
#             num_beams=4,  # You can adjust this to use beam search
#             early_stopping=True
#         )

#         # Decode the generated summary
#         summary = tokenizer.decode(summary_ids[0], skip_special_tokens=True)

#         # Print the generated summary
#         print("Generated Summary:", summary)


#     return render_template('algorithm_selection.html', summary=summary, error=None)






app.secret_key = 'your_secret_key'

# Set up Google API key
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
# Assuming genai is configured properly
import google.generativeai as genai
genai.configure(api_key=GOOGLE_API_KEY)

# Initialize Gemini model
model = genai.GenerativeModel('gemini-pro')

def generate_mcqs(text, number, tone):
    prompt = f"""
    Text: {text}
    
    You are an expert MCQ maker. Given the above text, create a quiz of {number} multiple choice questions in {tone} tone.
    Format like the example:
    
    1: Question?
    a) Option 1
    b) Option 2
    c) Option 3
    d) Option 4
    Correct answer: x
    """
    
    for attempt in range(3):  # Retry up to 3 times
        try:
            response = model.generate_content(prompt)
            quiz = response.text
            
            # Parse generated quiz
            mcqs = quiz.split('\n\n')
            quiz_data = []
            answer_key = {}
            for mcq in mcqs:
                lines = mcq.split('\n')
                if len(lines) < 5:
                    continue
                question = lines[0].strip()
                options = {chr(ord('a') + i): line[3:].strip() for i, line in enumerate(lines[1:5])}
                correct = lines[5][14:].strip()
                quiz_data.append({"Question": question, "Options": options})
                answer_key[question] = correct
                
            return quiz_data, answer_key
        
        except Exception as e:
            if "500" in str(e):
                time.sleep(2 ** attempt)  # Exponential backoff
            else:
                return None, None
    
    return None, None

def get_webpage_content(url):
    response = requests.get(url)
    soup = BeautifulSoup(response.content, 'html.parser')
    return soup.get_text()

@app.route('/prediction', methods=['GET', 'POST'])
def prediction():
    if request.method == "POST":
        text = None
        file = request.files["file"]
        webpage_url = request.form["webpage_url"]
        txt = request.form["txt"]
        number_of_questions = int(request.form["number_of_questions"])
        tone = request.form["tone"]

        
        # File or URL or Text Area content retrieval
        if file:
            text = file.read().decode("utf-8")
        elif webpage_url:
            text = get_webpage_content(webpage_url)
        elif txt:
            text = txt

        # Generate MCQs
        quiz_data, answer_key = generate_mcqs(text, number_of_questions, tone)
        
        if quiz_data:
            session['quiz_data'] = quiz_data
            session['answer_key'] = answer_key
            return render_template("prediction.html", quiz_data=quiz_data)
        else:
            flash("Failed to generate MCQs.")
            return redirect(url_for("index"))

    return render_template("prediction.html")




@app.route("/download/<format>")
def download(format):
    quiz_data = session.get("quiz_data", [])
    answer_key = session.get("answer_key", {})

    if format == "pdf":
        pdf_file = SimpleDocTemplate('mcqs.pdf', pagesize=letter)
        elements = []
        styles = getSampleStyleSheet()

        for mcq in quiz_data:
            elements.append(Paragraph(mcq['Question'], styles['Heading1']))
            for option, text in mcq['Options'].items():
                elements.append(Paragraph(f"{option}) {text}", styles['Normal']))
            elements.append(Paragraph("", styles['Normal']))

        pdf_file.build(elements)
        return send_file("mcqs.pdf", as_attachment=True)

    elif format == "word":
        document = Document()
        for mcq in quiz_data:
            document.add_heading(mcq['Question'], level=1)
            for option, text in mcq['Options'].items():
                document.add_paragraph(f"{option}) {text}")
            document.add_paragraph()

        document.save("mcqs.docx")
        return send_file("mcqs.docx", as_attachment=True)

    elif format == "answer_key":
        answer_key_doc = Document()
        answer_key_doc.add_heading("Answer Key", level=1)
        for question, answer in answer_key.items():
            answer_key_doc.add_paragraph(f"{question} - Correct answer: {answer}")

        answer_key_doc.save("answer_key.docx")
        return send_file("answer_key.docx", as_attachment=True)

    return redirect(url_for("results"))



if __name__ == '__main__':
    app.run(debug = True)